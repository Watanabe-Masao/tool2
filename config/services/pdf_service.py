#!/usr/bin/env python3
# -*- coding: utf-8 -*-

"""
PDF変換サービス

設計原則:
- 責務の明確化: PDF変換のみを担当
- 抽象化・隠蔽: LibreOfficeの呼び出し詳細を隠蔽
- 信頼性: エラーハンドリングとタイムアウト管理
"""

import logging
from pathlib import Path
import subprocess
import openpyxl
from datetime import datetime
from config.config import settings

# ロガー設定
logger = logging.getLogger(__name__)


class PDFService:
    """
    PDF変換サービス

    LibreOfficeを使用したExcelファイルのPDF変換を提供します。
    """

    @staticmethod
    def prepare_for_conversion(excel_path: Path) -> None:
        """
        PDF変換用にExcelファイルを最適化

        LibreOfficeのPDF変換で問題を起こす要素を除去：
        - 非表示列の内容をクリア（削除ではなく）
        - 列幅を極小値に設定
        - 日付を文字列に変換（Safari対応）

        Args:
            excel_path: 変換前のExcelファイルパス（このファイルを直接修正）

        設計原則:
        - 信頼性: エラー発生時も処理を続行（元のファイルを使用）
        """
        try:
            wb = openpyxl.load_workbook(excel_path)
            ws = wb.active

            # 非表示列の処理（A, Fのみ。AWは表示列なので対象外）
            PDFService._clear_hidden_columns(ws)

            # 日付を文字列に変換（Safari対応）
            PDFService._convert_dates_to_strings(ws)

            # 変更を保存
            wb.save(excel_path)
            wb.close()
            logger.debug("Excel file optimized for PDF conversion")

        except Exception as e:
            logger.error(f"PDF conversion preparation failed: {e}", exc_info=True)
            # エラー発生時も処理を続行（元のファイルを使用）
            try:
                wb.close()
            except:
                pass

    @staticmethod
    def _clear_hidden_columns(ws) -> None:
        """
        非表示列の内容をクリア

        Args:
            ws: openpyxlワークシート
        """
        hidden_columns = ['A', 'F']

        for col_letter in hidden_columns:
            if ws.column_dimensions[col_letter].hidden:
                logger.debug(f"Processing hidden column {col_letter} for PDF conversion")

                # 列の全セルの内容をクリア（結合セルはスキップ）
                for row in range(1, ws.max_row + 1):
                    cell = ws[f'{col_letter}{row}']
                    # 結合セルの場合はスキップ
                    if not isinstance(cell, openpyxl.cell.cell.MergedCell):
                        cell.value = None
                        cell.number_format = 'General'

                # 列幅を極小値に設定（非表示のまま維持）
                ws.column_dimensions[col_letter].width = 0.08333
                logger.debug(f"Column {col_letter}: cleared, width=0.08333")

    @staticmethod
    def _convert_dates_to_strings(ws) -> None:
        """
        日付セルをPDF変換に適した形式に変換（Safari対応）

        Args:
            ws: openpyxlワークシート
        """
        # 日本語曜日マッピング
        weekday_ja = ['月', '火', '水', '木', '金', '土', '日']

        for row in ws.iter_rows(min_row=7, max_row=100):  # データエリア
            for cell in row:
                if isinstance(cell.value, datetime):
                    # datetimeを読みやすい文字列に変換（日本語曜日付き）
                    weekday_str = weekday_ja[cell.value.weekday()]
                    date_str = cell.value.strftime(f'%m/%d({weekday_str})')
                    cell.value = date_str
                    cell.number_format = '@'  # テキスト形式
                    logger.debug(f"Converted datetime in {cell.coordinate} to string: {date_str}")

    @staticmethod
    def convert_to_pdf(excel_path: Path, pdf_path: Path) -> bool:
        """
        ExcelファイルをLibreOfficeを使ってPDFに変換（完全な書式保持）

        Excel結合セルのPDF変換対策：Excel → ODS → PDF の2段階変換
        理由：LibreOfficeがExcelを直接PDFに変換すると、結合セルの値が失われることがある
              ODSを経由することで、LibreOffice自身が結合セルを正しく解釈する

        Args:
            excel_path: Excelファイルのパス
            pdf_path: 出力PDFファイルのパス

        Returns:
            bool: 変換成功時True

        Raises:
            Exception: 変換失敗時

        設計原則:
        - 抽象化: LibreOffice呼び出しの詳細を隠蔽
        - 信頼性: タイムアウト管理とエラーハンドリング
        """
        try:
            # Step 1: Excel → ODS
            ods_path = excel_path.parent / f"{excel_path.stem}.ods"
            result_ods = subprocess.run(
                [
                    'libreoffice',
                    '--headless',
                    '--convert-to', 'ods',
                    '--outdir', str(excel_path.parent),
                    str(excel_path)
                ],
                capture_output=True,
                text=True,
                timeout=settings.libreoffice_timeout
            )

            if result_ods.returncode != 0:
                error_msg = f"Excel to ODS conversion failed:\nstdout: {result_ods.stdout}\nstderr: {result_ods.stderr}"
                logger.error(error_msg)
                raise Exception(error_msg)

            # Step 2: ODS → PDF
            result_pdf = subprocess.run(
                [
                    'libreoffice',
                    '--headless',
                    '--convert-to', 'pdf',
                    '--outdir', str(pdf_path.parent),
                    str(ods_path)
                ],
                capture_output=True,
                text=True,
                timeout=settings.libreoffice_timeout
            )

            # ODS一時ファイルを削除
            if ods_path.exists():
                ods_path.unlink()

            if result_pdf.returncode != 0:
                error_msg = f"ODS to PDF conversion failed:\nstdout: {result_pdf.stdout}\nstderr: {result_pdf.stderr}"
                logger.error(error_msg)
                raise Exception(error_msg)

            # 出力ファイル名を調整（LibreOfficeは元のファイル名で出力する）
            expected_pdf = pdf_path.parent / f"{excel_path.stem}.pdf"
            if expected_pdf.exists() and expected_pdf != pdf_path:
                expected_pdf.rename(pdf_path)

            return pdf_path.exists()

        except subprocess.TimeoutExpired:
            raise Exception(f"PDF変換がタイムアウトしました（{settings.libreoffice_timeout}秒）")
        except Exception as e:
            raise Exception(f"PDF変換中にエラーが発生: {str(e)}")
