#!/usr/bin/env python3
# -*- coding: utf-8 -*-

"""
配分表テンプレート作成WebアプリケーションAPI
FastAPIを使用したバックエンドサーバー
"""

import os
import io
import uuid
from datetime import datetime
from typing import Optional
from pathlib import Path
from urllib.parse import quote

from fastapi import FastAPI, HTTPException, Response
from fastapi.staticfiles import StaticFiles
from fastapi.templating import Jinja2Templates
from fastapi.responses import StreamingResponse, HTMLResponse, FileResponse
from fastapi import Request
from pydantic import BaseModel, Field
from typing import Dict, List

from haibun_template_creator import (
    HaibunTemplateCreator,
    TemplateConfig,
    StoreData,
    ProductData
)

import openpyxl
import subprocess


# アプリケーションバージョン（静的ファイルのキャッシュバスティング用）
APP_VERSION = "2.2.1"

# FastAPIアプリケーション初期化
app = FastAPI(
    title="配分表テンプレート作成API",
    description="Excelの配分表テンプレートを生成するWebアプリケーション",
    version=APP_VERSION
)

# 静的ファイルとテンプレートの設定
app.mount("/static", StaticFiles(directory="static"), name="static")
templates = Jinja2Templates(directory="templates")


# キャッシュ制御ミドルウェア
@app.middleware("http")
async def add_cache_control_header(request: Request, call_next):
    """
    静的ファイルにCache-Controlヘッダーを追加
    バージョンクエリパラメータがある場合は長期間キャッシュ
    """
    response = await call_next(request)

    # 静的ファイルの場合
    if request.url.path.startswith("/static/"):
        # バージョンパラメータがある場合は1年間キャッシュ
        if "v=" in request.url.query:
            response.headers["Cache-Control"] = "public, max-age=31536000, immutable"
        else:
            # バージョンパラメータがない場合は短期間のみキャッシュ
            response.headers["Cache-Control"] = "public, max-age=3600"

    return response

# 一時ファイル保存ディレクトリ
TEMP_DIR = Path("temp_files")
TEMP_DIR.mkdir(exist_ok=True)


def prepare_excel_for_pdf_conversion(excel_path: Path) -> None:
    """
    PDF変換用にExcelファイルを最適化

    非表示列を物理的に削除してLibreOfficeのPDF変換での列ずれを防止：
    - 非表示列（A, F）を完全に削除
    - 日付を文字列に変換（Safari対応）

    Args:
        excel_path: 変換前のExcelファイルパス（このファイルを直接修正）
    """
    import openpyxl
    from datetime import datetime
    from copy import copy

    try:
        wb = openpyxl.load_workbook(excel_path)
        ws = wb.active

        print(f"[DEBUG] Starting PDF optimization...")

        # ステップ1: 削除対象の非表示列を特定
        columns_to_delete = []

        # 非表示列をチェック（列インデックスで記録）
        # 後ろから削除するため、大きい順
        hidden_specs = [
            ('F', 6),
            ('A', 1),
        ]

        for col_letter, col_idx in hidden_specs:
            if col_letter in ws.column_dimensions and ws.column_dimensions[col_letter].hidden:
                columns_to_delete.append((col_idx, col_letter))
                print(f"[DEBUG] Marking column {col_letter} (index {col_idx}) for deletion")

        # ステップ2: 後ろから順に列を削除（インデックスのずれを防ぐ）
        for col_idx, col_letter in sorted(columns_to_delete, reverse=True):
            try:
                print(f"[DEBUG] Deleting column {col_letter} at index {col_idx}")
                ws.delete_cols(col_idx, 1)
                print(f"[DEBUG] Successfully deleted column {col_letter}")
            except Exception as e:
                print(f"[WARNING] Failed to delete column {col_letter}: {e}")
                # 削除失敗の場合は内容をクリア
                for row in range(1, ws.max_row + 1):
                    try:
                        cell = ws.cell(row=row, column=col_idx)
                        if not isinstance(cell, openpyxl.cell.cell.MergedCell):
                            cell.value = None
                    except:
                        pass
                if col_letter in ws.column_dimensions:
                    ws.column_dimensions[col_letter].width = 0.001

        # ステップ3: 日付セルを文字列に変換（Safari対応）
        weekday_ja = ['月', '火', '水', '木', '金', '土', '日']

        for row in ws.iter_rows(min_row=1, max_row=100):
            for cell in row:
                if cell.value and isinstance(cell.value, datetime):
                    weekday_str = weekday_ja[cell.value.weekday()]
                    date_str = cell.value.strftime(f'%m/%d({weekday_str})')
                    cell.value = date_str
                    cell.number_format = '@'
                    print(f"[DEBUG] Converted datetime in {cell.coordinate} to: {date_str}")

        # ステップ4: ファイルを保存
        wb.save(excel_path)
        wb.close()
        print(f"[DEBUG] Excel file optimized for PDF conversion")

    except Exception as e:
        print(f"[ERROR] PDF conversion preparation failed: {e}")
        import traceback
        print(f"[ERROR] Traceback:\n{traceback.format_exc()}")
        # エラー発生時も処理を続行（元のファイルを使用）
        try:
            wb.close()
        except:
            pass


def excel_to_pdf(excel_path: Path, pdf_path: Path) -> bool:
    """
    ExcelファイルをLibreOfficeを使ってPDFに変換（完全な書式保持）

    Args:
        excel_path: Excelファイルのパス
        pdf_path: 出力PDFファイルのパス

    Returns:
        bool: 変換成功時True

    Raises:
        Exception: 変換失敗時
    """
    try:
        # Excel結合セルのPDF変換対策：Excel → ODS → PDF の2段階変換
        # 理由：LibreOfficeがExcelを直接PDFに変換すると、結合セルの値が失われることがある
        #       ODSを経由することで、LibreOffice自身が結合セルを正しく解釈する

        # Step 1: Excel → ODS
        ods_path = excel_path.parent / f"{excel_path.stem}.ods"
        result_ods = subprocess.run(
            [
                'libreoffice',
                '--headless',
                '--convert-to', 'ods',
                '--outdir', str(excel_path.parent),
                str(excel_path)
            ],
            capture_output=True,
            text=True,
            timeout=30
        )

        if result_ods.returncode != 0:
            error_msg = f"Excel to ODS conversion failed:\nstdout: {result_ods.stdout}\nstderr: {result_ods.stderr}"
            print(f"[ERROR] {error_msg}")
            raise Exception(error_msg)

        # Step 2: ODS → PDF
        result_pdf = subprocess.run(
            [
                'libreoffice',
                '--headless',
                '--convert-to', 'pdf',
                '--outdir', str(pdf_path.parent),
                str(ods_path)
            ],
            capture_output=True,
            text=True,
            timeout=30
        )

        # ODS一時ファイルを削除
        if ods_path.exists():
            ods_path.unlink()

        if result_pdf.returncode != 0:
            error_msg = f"ODS to PDF conversion failed:\nstdout: {result_pdf.stdout}\nstderr: {result_pdf.stderr}"
            print(f"[ERROR] {error_msg}")
            raise Exception(error_msg)

        # 出力ファイル名を調整（LibreOfficeは元のファイル名で出力する）
        expected_pdf = pdf_path.parent / f"{excel_path.stem}.pdf"
        if expected_pdf.exists() and expected_pdf != pdf_path:
            expected_pdf.rename(pdf_path)

        return pdf_path.exists()

    except subprocess.TimeoutExpired:
        raise Exception("PDF変換がタイムアウトしました（30秒）")
    except Exception as e:
        raise Exception(f"PDF変換中にエラーが発生: {str(e)}")


# リクエストモデル
class ProductDataRequest(BaseModel):
    """商品データリクエスト"""
    name: Optional[str] = Field(default=None, max_length=50, description="品名")
    product_name: Optional[str] = Field(default=None, max_length=50, description="品名（後方互換用）")
    delivery_date: Optional[str] = Field(default=None, description="納品日（YYYY-MM-DD形式）")
    origin: Optional[str] = Field(default=None, max_length=30, description="産地")
    standard: Optional[str] = Field(default=None, max_length=20, description="規格")
    store_cost: Optional[float] = Field(default=None, description="店着原価")
    price: Optional[float] = Field(default=None, description="税抜売価")
    quantity: Optional[int] = Field(default=None, description="入数")
    total_delivery: Optional[int] = Field(default=None, description="総納品数")
    delivery_dest: Optional[str] = Field(default=None, max_length=30, description="納品先")
    store_quantities: Dict[str, int] = Field(default_factory=dict, description="店舗配分数")


class TemplateRequest(BaseModel):
    """テンプレート生成リクエスト（Phase 3: 5-step workflow対応）"""
    # Step 1: 店着日（全商品共通）
    delivery_date: Optional[str] = Field(default=None, description="店着日（YYYY-MM-DD形式）")
    # Step 2: 帳合先（全商品共通）
    supplier: Optional[str] = Field(default=None, max_length=50, description="帳合先名")
    # 商品数（自動計算されるが、後方互換用に残す）
    num_blocks: Optional[int] = Field(default=None, ge=1, le=100, description="商品ブロック数（1-100）")
    output_filename: Optional[str] = Field(
        default=None,
        description="出力ファイル名（省略時は自動生成）"
    )
    buyer_name: Optional[str] = Field(
        default=None,
        max_length=20,
        description="担当バイヤー名（最大20文字）"
    )
    pixel_100: Optional[float] = Field(default=13.5714285714, description="100ピクセル列幅")
    pixel_50: Optional[float] = Field(default=6.4285714286, description="50ピクセル列幅")
    # Step 3-5: 商品情報（動的リスト）
    products: List[ProductDataRequest] = Field(default_factory=list, description="商品データリスト")


# レスポンスモデル
class TemplateResponse(BaseModel):
    """テンプレート生成レスポンス"""
    success: bool
    message: str
    download_url: Optional[str] = None
    filename: Optional[str] = None


@app.get("/api/health")
async def health_check():
    """
    ヘルスチェックエンドポイント
    Renderがサーバーの状態を確認するために使用
    """
    return {
        "status": "healthy",
        "version": APP_VERSION,
        "service": "tool2"
    }


@app.get("/", response_class=HTMLResponse)
async def root(request: Request):
    """
    ルートページ - フロントエンドUIを表示
    """
    response = templates.TemplateResponse("index.html", {
        "request": request,
        "version": APP_VERSION
    })
    # HTMLページはキャッシュしない（常に最新版を取得）
    response.headers["Cache-Control"] = "no-cache, no-store, must-revalidate"
    response.headers["Pragma"] = "no-cache"
    response.headers["Expires"] = "0"
    return response


@app.get("/login", response_class=HTMLResponse)
async def login_page(request: Request):
    """
    ログインページ
    """
    response = templates.TemplateResponse("login.html", {
        "request": request
    })
    response.headers["Cache-Control"] = "no-cache, no-store, must-revalidate"
    response.headers["Pragma"] = "no-cache"
    response.headers["Expires"] = "0"
    return response


@app.post("/api/generate", response_model=TemplateResponse)
async def generate_template(req: TemplateRequest):
    """
    テンプレート生成API

    Args:
        req: テンプレート生成リクエスト

    Returns:
        TemplateResponse: 生成結果とダウンロードURL
    """
    try:
        # ファイル名の生成
        if req.output_filename:
            filename = req.output_filename
            if not filename.endswith('.xlsx'):
                filename += '.xlsx'
        else:
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            filename = f"配分表_テンプレート_{timestamp}.xlsx"

        # 一時ファイルパス
        file_id = str(uuid.uuid4())
        temp_path = TEMP_DIR / f"{file_id}.xlsx"

        # 商品数を動的に取得（num_blocksが指定されていない場合）
        num_blocks = req.num_blocks if req.num_blocks else len(req.products)

        # 設定作成
        config = TemplateConfig(
            num_blocks=num_blocks,
            pixel_100=req.pixel_100,
            pixel_50=req.pixel_50,
            default_output_path=str(temp_path)
        )

        # 商品データをProductDataオブジェクトに変換
        products = []
        for product_req in req.products:
            # product_nameの決定（name優先、なければproduct_nameにフォールバック）
            product_name = product_req.name or product_req.product_name

            # delivery_dateの決定（商品固有 > 全体共通）
            delivery_date = product_req.delivery_date or req.delivery_date

            product_data = ProductData(
                delivery_date=delivery_date,
                origin=product_req.origin,
                standard=product_req.standard,
                product_name=product_name,
                store_cost=product_req.store_cost,
                price=product_req.price,
                quantity=product_req.quantity,
                total_delivery=product_req.total_delivery,
                delivery_dest=product_req.delivery_dest or req.supplier,  # 納品先がなければ帳合先を使用
                store_quantities=product_req.store_quantities
            )
            products.append(product_data)

        # テンプレート生成
        creator = HaibunTemplateCreator(config=config)
        output_path = creator.create_template(
            buyer_name=req.buyer_name,
            products=products
        )

        # ダウンロードURL生成
        download_url = f"/api/download/{file_id}?filename={filename}"

        return TemplateResponse(
            success=True,
            message="テンプレートの生成に成功しました",
            download_url=download_url,
            filename=filename
        )

    except Exception as e:
        raise HTTPException(
            status_code=500,
            detail=f"テンプレート生成中にエラーが発生しました: {str(e)}"
        )


@app.get("/api/download/{file_id}")
async def download_template(file_id: str, filename: str = "配分表_テンプレート.xlsx"):
    """
    生成されたテンプレートファイルをダウンロード

    Args:
        file_id: ファイルID
        filename: ダウンロード時のファイル名

    Returns:
        StreamingResponse: Excelファイル
    """
    try:
        # ファイルパス取得
        temp_path = TEMP_DIR / f"{file_id}.xlsx"

        if not temp_path.exists():
            raise HTTPException(
                status_code=404,
                detail="ファイルが見つかりません"
            )

        # ファイルを読み込み
        with open(temp_path, "rb") as f:
            file_content = f.read()

        # 注意: ファイルは削除せず、複数回ダウンロード可能にする
        # クリーンアップはshutdownイベントで実行される

        # 日本語ファイル名のエンコード（RFC 5987対応）
        # ASCIIフォールバック用にtemplate.xlsxを設定
        encoded_filename = quote(filename.encode('utf-8'))

        # ストリーミングレスポンス
        return StreamingResponse(
            io.BytesIO(file_content),
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={
                # RFC 5987形式: ASCIIフォールバック + UTF-8エンコード
                "Content-Disposition": f"attachment; filename=\"template.xlsx\"; filename*=UTF-8''{encoded_filename}"
            }
        )

    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(
            status_code=500,
            detail=f"ファイルのダウンロード中にエラーが発生しました: {str(e)}"
        )


@app.post("/api/preview")
async def preview_template(req: TemplateRequest):
    """
    テンプレートのPDFプレビューを生成

    Args:
        req: テンプレート生成リクエスト

    Returns:
        FileResponse: PDFファイル
    """
    import traceback

    try:
        # 一時ファイルパス
        file_id = str(uuid.uuid4())
        temp_excel_path = TEMP_DIR / f"{file_id}.xlsx"
        temp_pdf_path = TEMP_DIR / f"{file_id}.pdf"

        # 商品数を取得
        num_blocks = len(req.products) if req.products else 1

        # 設定作成
        config = TemplateConfig(
            num_blocks=num_blocks,
            pixel_100=req.pixel_100,
            pixel_50=req.pixel_50,
            default_output_path=str(temp_excel_path)
        )

        # 商品データをProductDataオブジェクトに変換
        products = []
        for product_req in req.products:
            # product_nameの決定（name優先、なければproduct_nameにフォールバック）
            product_name = product_req.name or product_req.product_name

            # delivery_dateの決定（商品固有 > 全体共通）
            delivery_date = product_req.delivery_date or req.delivery_date

            product_data = ProductData(
                delivery_date=delivery_date,
                origin=product_req.origin,
                standard=product_req.standard,
                product_name=product_name,
                store_cost=product_req.store_cost,
                price=product_req.price,
                quantity=product_req.quantity,
                total_delivery=product_req.total_delivery,
                delivery_dest=product_req.delivery_dest or req.supplier,  # 納品先がなければ帳合先を使用
                store_quantities=product_req.store_quantities
            )
            products.append(product_data)

        # テンプレート生成
        print(f"[DEBUG] Creating Excel template...")
        creator = HaibunTemplateCreator(config=config)
        creator.create_template(
            buyer_name=req.buyer_name,
            products=products
        )
        print(f"[DEBUG] Excel template created: {temp_excel_path.exists()}")

        # PDF変換用にExcelを最適化（非表示列を削除）
        print(f"[DEBUG] Preparing Excel for PDF conversion...")
        prepare_excel_for_pdf_conversion(temp_excel_path)

        # ExcelをPDFに変換（LibreOffice使用）
        print(f"[DEBUG] Converting Excel to PDF using LibreOffice...")
        excel_to_pdf(temp_excel_path, temp_pdf_path)
        print(f"[DEBUG] PDF created: {temp_pdf_path.exists()}")

        if not temp_pdf_path.exists():
            raise Exception("PDFファイルが生成されませんでした")

        # PDFを返す
        return FileResponse(
            path=str(temp_pdf_path),
            media_type="application/pdf",
            filename="preview.pdf",
            headers={
                "Content-Disposition": "inline; filename=preview.pdf"
            }
        )

    except Exception as e:
        error_detail = f"{str(e)}\n\nTraceback:\n{traceback.format_exc()}"
        print(f"[ERROR] PDF preview generation failed:\n{error_detail}")
        raise HTTPException(
            status_code=500,
            detail=f"PDFプレビュー生成中にエラーが発生しました: {str(e)}"
        )


@app.get("/api/health")
@app.head("/api/health")
async def health_check():
    """
    ヘルスチェックエンドポイント（GET/HEADメソッド対応）
    """
    return {"status": "ok", "message": "API is running"}


@app.get("/api/version")
async def get_version():
    """
    現在のアプリケーションバージョンを返す
    """
    return {
        "version": APP_VERSION,
        "pdf_preview_available": True,
        "cache_busting_enabled": True
    }


@app.get("/api/firebase-config")
async def get_firebase_config():
    """
    Firebaseの設定を環境変数から取得して返す
    セキュリティ向上のため、クライアント側にハードコードしない

    注意: Firebase Web SDKの仕様上、これらの設定は公開されても問題ありません。
    セキュリティはFirestore Security Rulesで制御します。
    """
    firebase_config = {
        "apiKey": os.getenv("FIREBASE_API_KEY", "AIzaSyCjuPCpB0wqHxdX4JWL6VnEj1LJWgr4cKc"),
        "authDomain": os.getenv("FIREBASE_AUTH_DOMAIN", "haibun-distribution.firebaseapp.com"),
        "projectId": os.getenv("FIREBASE_PROJECT_ID", "haibun-distribution"),
        "storageBucket": os.getenv("FIREBASE_STORAGE_BUCKET", "haibun-distribution.firebasestorage.app"),
        "messagingSenderId": os.getenv("FIREBASE_MESSAGING_SENDER_ID", "742220611313"),
        "appId": os.getenv("FIREBASE_APP_ID", "1:742220611313:web:bec0f006c4c648adcbb350")
    }

    return firebase_config


@app.head("/")
async def root_head():
    """
    ルートパスのHEADリクエスト対応（Renderヘルスチェック用）
    """
    return Response(status_code=200)


@app.on_event("startup")
async def startup_event():
    """
    アプリケーション起動時の処理
    """
    port = os.getenv("PORT", "8000")
    print("=" * 60)
    print("配分表テンプレート作成Webアプリケーションを起動しました")
    print(f"PORT: {port}")
    print("=" * 60)


@app.on_event("shutdown")
async def shutdown_event():
    """
    アプリケーション終了時の処理 - 一時ファイルのクリーンアップ
    """
    print("一時ファイルをクリーンアップ中...")
    try:
        for file in TEMP_DIR.glob("*.xlsx"):
            file.unlink()
        for file in TEMP_DIR.glob("*.pdf"):
            file.unlink()
        print("クリーンアップ完了")
    except Exception as e:
        print(f"クリーンアップ中にエラーが発生: {e}")


if __name__ == "__main__":
    import uvicorn
    port = int(os.getenv("PORT", 8000))
    uvicorn.run(app, host="0.0.0.0", port=port)
