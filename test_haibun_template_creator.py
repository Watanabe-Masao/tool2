#!/usr/bin/env python3
# -*- coding: utf-8 -*-

"""
配分表テンプレート作成のテストコード

ヘッダー構造、データ配置、書式設定などを検証して
将来の変更でレイアウトが崩れないことを保証します。
"""

import pytest
import openpyxl
from pathlib import Path
from datetime import datetime
from haibun_template_creator import (
    HaibunTemplateCreator,
    TemplateConfig,
    ProductData
)


@pytest.fixture
def temp_output_path(tmp_path):
    """一時出力パスを生成"""
    return tmp_path / "test_template.xlsx"


@pytest.fixture
def sample_product_data():
    """サンプル商品データを生成"""
    return ProductData(
        delivery_date='2025-01-15',
        origin='高知県',
        standard='L',
        product_name='トマト',
        store_cost=120.5,
        price=198.0,
        quantity=10,
        total_delivery=100,
        delivery_dest='本社',
        store_quantities={'01': 10, '02': 15, '03': 20}
    )


class TestHeaderStructure:
    """ヘッダー構造のテスト（7-8行目）"""

    def test_header_cell_merges(self, temp_output_path):
        """ヘッダーのセル結合が正しいかテスト"""
        config = TemplateConfig(num_blocks=1, default_output_path=str(temp_output_path))
        creator = HaibunTemplateCreator(config=config)
        creator.create_template()

        wb = openpyxl.load_workbook(temp_output_path)
        ws = wb.active

        # 結合されているセル範囲を確認
        merged_ranges = [str(merged_range) for merged_range in ws.merged_cells.ranges]

        # 期待される結合
        assert 'A7:A8' in merged_ranges, "A7:A8が結合されていません（商品コード）"
        assert 'B7:C8' in merged_ranges, "B7:C8が結合されていません（納品日）"
        assert 'D8:E8' in merged_ranges, "D8:E8が結合されていません（品名）"
        assert 'F7:F8' in merged_ranges, "F7:F8が結合されていません（LFC着）"
        assert 'G7:G8' in merged_ranges, "G7:G8が結合されていません（店着原価）"
        assert 'AT7:AT8' in merged_ranges, "AT7:AT8が結合されていません（合計）"
        assert 'AV7:AV8' in merged_ranges, "AV7:AV8が結合されていません（納品先）"

        wb.close()

    def test_header_text_positions(self, temp_output_path):
        """ヘッダーテキストが正しい位置にあるかテスト"""
        config = TemplateConfig(num_blocks=1, default_output_path=str(temp_output_path))
        creator = HaibunTemplateCreator(config=config)
        creator.create_template()

        wb = openpyxl.load_workbook(temp_output_path)
        ws = wb.active

        # 7行目のヘッダー
        assert ws['A7'].value == '商品コード', "A7のヘッダーが正しくありません"
        assert ws['B7'].value == '納品日', "B7のヘッダーが正しくありません"
        assert ws['D7'].value == '産地', "D7のヘッダーが正しくありません"
        assert ws['E7'].value == '規  格', "E7のヘッダーが正しくありません"
        assert ws['F7'].value == ' LFC着', "F7のヘッダーが正しくありません"
        assert ws['G7'].value == '店着原価', "G7のヘッダーが正しくありません"
        assert ws['H7'].value == '税抜', "H7のヘッダーが正しくありません"
        assert ws['I7'].value == 'ｹｰｽ', "I7のヘッダーが正しくありません"

        # 8行目のヘッダー（重要：H8とI8の位置）
        assert ws['D8'].value == '品 名', "D8のヘッダーが正しくありません"
        assert ws['H8'].value == '税込', "H8のヘッダーが正しくありません（税抜の下にあるべき）"
        assert ws['I8'].value == '入数', "I8のヘッダーが正しくありません（ケースの下にあるべき）"
        assert ws['AU7'].value == '納品数', "AU7のヘッダーが正しくありません"
        assert ws['AU8'].value == '差異', "AU8のヘッダーが正しくありません"

        # 列インデックスの厳密な確認
        assert ws.cell(7, 8).value == '税抜', "H7 (col=8) が税抜ではありません"
        assert ws.cell(8, 8).value == '税込', "H8 (col=8) が税込ではありません"
        assert ws.cell(7, 9).value == 'ｹｰｽ', "I7 (col=9) がケースではありません"
        assert ws.cell(8, 9).value == '入数', "I8 (col=9) が入数ではありません"

        wb.close()

    def test_store_headers(self, temp_output_path):
        """店舗ヘッダー（J-AS列）が正しいかテスト"""
        config = TemplateConfig(num_blocks=1, default_output_path=str(temp_output_path))
        creator = HaibunTemplateCreator(config=config)
        creator.create_template()

        wb = openpyxl.load_workbook(temp_output_path)
        ws = wb.active

        # 店舗コードの確認（一部サンプル）
        assert ws['J7'].value in ['01', 1], "J7の店舗コードが正しくありません"
        assert ws['K7'].value in ['02', 2], "K7の店舗コードが正しくありません"
        assert ws['AS7'].value == 911, "AS7の店舗コードが正しくありません（惣菜）"

        # 店舗名の確認（一部サンプル）
        assert ws['J8'].value == '朝倉', "J8の店舗名が正しくありません"
        assert ws['K8'].value == '伊野', "K8の店舗名が正しくありません"
        assert ws['AS8'].value == '惣菜', "AS8の店舗名が正しくありません"

        wb.close()


class TestBasicFunctionality:
    """基本機能のテスト"""

    def test_buyer_name_input(self, temp_output_path):
        """担当バイヤー名が正しく入力されるかテスト"""
        buyer_name = "山田太郎"
        config = TemplateConfig(num_blocks=1, default_output_path=str(temp_output_path))
        creator = HaibunTemplateCreator(config=config)
        creator.create_template(buyer_name=buyer_name)

        wb = openpyxl.load_workbook(temp_output_path)
        ws = wb.active

        # AE4:AT4セルに担当バイヤー名が入力されているか
        assert ws['AE4'].value == buyer_name, "担当バイヤー名が正しく入力されていません"

        wb.close()

    def test_product_data_input(self, temp_output_path, sample_product_data):
        """商品データが正しく入力されるかテスト"""
        config = TemplateConfig(num_blocks=1, default_output_path=str(temp_output_path))
        creator = HaibunTemplateCreator(config=config)
        creator.create_template(products=[sample_product_data])

        wb = openpyxl.load_workbook(temp_output_path)
        ws = wb.active

        # 商品データが正しいセルに入力されているか（1商品目は9行目から）
        # 納品日（B9）- datetimeオブジェクトとして格納される
        assert ws['B9'].value is not None, "納品日が入力されていません"

        # 産地（D9）
        assert ws['D9'].value == '高知県', "産地が正しく入力されていません"

        # 規格（E9）
        assert ws['E9'].value == 'L', "規格が正しく入力されていません"

        # 品名（D10）- detail_row
        assert ws['D10'].value == 'トマト', "品名が正しく入力されていません"

        # 店着原価（G9）
        assert ws['G9'].value == 120.5, "店着原価が正しく入力されていません"

        # 税抜売価（H9）
        assert ws['H9'].value == 198.0, "税抜売価が正しく入力されていません"

        # 入数（I10）- detail_row
        assert ws['I10'].value == 10, "入数が正しく入力されていません"

        # 総納品数（AU9）
        assert ws['AU9'].value == 100, "総納品数が正しく入力されていません"

        # 納品先（AV9）
        assert ws['AV9'].value == '本社', "納品先が正しく入力されていません"

        wb.close()

    def test_store_quantities_input(self, temp_output_path, sample_product_data):
        """店舗配分数が正しい列に入力されるかテスト"""
        config = TemplateConfig(num_blocks=1, default_output_path=str(temp_output_path))
        creator = HaibunTemplateCreator(config=config)
        creator.create_template(products=[sample_product_data])

        wb = openpyxl.load_workbook(temp_output_path)
        ws = wb.active

        # 店舗配分数が正しい列に入力されているか（detail_row = 10行目）
        assert ws['J10'].value == 10, "店舗01の配分数が正しく入力されていません"
        assert ws['K10'].value == 15, "店舗02の配分数が正しく入力されていません"
        assert ws['L10'].value == 20, "店舗03の配分数が正しく入力されていません"

        wb.close()

    def test_multiple_products(self, temp_output_path):
        """複数商品のデータが正しく配置されるかテスト"""
        products = [
            ProductData(product_name='商品1', origin='産地1'),
            ProductData(product_name='商品2', origin='産地2'),
            ProductData(product_name='商品3', origin='産地3'),
        ]

        config = TemplateConfig(num_blocks=3, default_output_path=str(temp_output_path))
        creator = HaibunTemplateCreator(config=config)
        creator.create_template(products=products)

        wb = openpyxl.load_workbook(temp_output_path)
        ws = wb.active

        # 各商品が3行間隔で配置されているか
        # 商品1: 9行目（data_row）、10行目（detail_row）
        assert ws['D9'].value == '産地1', "商品1の産地が正しく配置されていません"
        assert ws['D10'].value == '商品1', "商品1の品名が正しく配置されていません"

        # 商品2: 12行目（data_row）、13行目（detail_row）
        assert ws['D12'].value == '産地2', "商品2の産地が正しく配置されていません"
        assert ws['D13'].value == '商品2', "商品2の品名が正しく配置されていません"

        # 商品3: 15行目（data_row）、16行目（detail_row）
        assert ws['D15'].value == '産地3', "商品3の産地が正しく配置されていません"
        assert ws['D16'].value == '商品3', "商品3の品名が正しく配置されていません"

        wb.close()


class TestFormatting:
    """書式設定のテスト"""

    def test_print_settings(self, temp_output_path):
        """印刷設定が正しいかテスト"""
        config = TemplateConfig(num_blocks=1, default_output_path=str(temp_output_path))
        creator = HaibunTemplateCreator(config=config)
        creator.create_template()

        wb = openpyxl.load_workbook(temp_output_path)
        ws = wb.active

        # 印刷設定の確認
        assert ws.page_setup.orientation == 'landscape', "印刷方向が横向きではありません"
        assert ws.page_setup.paperSize == 9, "用紙サイズがA4ではありません"
        # fitToPageはsheet_propertiesを経由してアクセス
        if ws.sheet_properties and ws.sheet_properties.pageSetUpPr:
            assert ws.sheet_properties.pageSetUpPr.fitToPage == True, "fitToPageが有効になっていません"
        assert ws.page_setup.fitToWidth == 1, "fitToWidthが1に設定されていません"
        assert ws.page_setup.fitToHeight == 0, "fitToHeightが0に設定されていません"

        wb.close()

    def test_column_widths(self, temp_output_path):
        """列幅が正しく設定されているかテスト"""
        config = TemplateConfig(num_blocks=1, default_output_path=str(temp_output_path))
        creator = HaibunTemplateCreator(config=config)
        creator.create_template()

        wb = openpyxl.load_workbook(temp_output_path)
        ws = wb.active

        # 非表示列の確認
        assert ws.column_dimensions['A'].hidden == True, "A列が非表示になっていません"
        assert ws.column_dimensions['F'].hidden == True, "F列が非表示になっていません"
        assert ws.column_dimensions['AW'].hidden == True, "AW列が非表示になっていません"

        # 一部の列幅を確認
        assert ws.column_dimensions['B'].width == 7.296875, "B列の幅が正しくありません"
        assert ws.column_dimensions['I'].width == 13.5714285714, "I列の幅が正しくありません"

        wb.close()

    def test_row_heights(self, temp_output_path):
        """行高が正しく設定されているかテスト"""
        config = TemplateConfig(num_blocks=1, default_output_path=str(temp_output_path))
        creator = HaibunTemplateCreator(config=config)
        creator.create_template()

        wb = openpyxl.load_workbook(temp_output_path)
        ws = wb.active

        # 行高の確認
        assert ws.row_dimensions[1].height == 30.0, "1行目の高さが正しくありません（40px相当）"
        assert ws.row_dimensions[2].hidden == True, "2行目が非表示になっていません"
        assert ws.row_dimensions[3].height == 22.5, "3行目の高さが正しくありません（30px相当）"
        assert ws.row_dimensions[4].height == 22.5, "4行目の高さが正しくありません（30px相当）"
        assert ws.row_dimensions[5].height == 18.75, "5行目の高さが正しくありません（25px相当）"
        assert ws.row_dimensions[6].height == 18.75, "6行目の高さが正しくありません（25px相当）"

        wb.close()


class TestFormulas:
    """数式のテスト"""

    def test_tax_formula(self, temp_output_path, sample_product_data):
        """税込価格の数式が正しいかテスト"""
        config = TemplateConfig(num_blocks=1, default_output_path=str(temp_output_path))
        creator = HaibunTemplateCreator(config=config)
        creator.create_template(products=[sample_product_data])

        wb = openpyxl.load_workbook(temp_output_path)
        ws = wb.active

        # H10（detail_row）に税込価格の数式があるか
        assert ws['H10'].value == '=H9*1.08', "税込価格の数式が正しくありません"

        wb.close()

    def test_total_formula(self, temp_output_path, sample_product_data):
        """合計の数式が正しいかテスト"""
        config = TemplateConfig(num_blocks=1, default_output_path=str(temp_output_path))
        creator = HaibunTemplateCreator(config=config)
        creator.create_template(products=[sample_product_data])

        wb = openpyxl.load_workbook(temp_output_path)
        ws = wb.active

        # AT9（data_row）に合計の数式があるか
        assert ws['AT9'].value == '=SUM(J10:AS11)', "合計の数式が正しくありません"

        wb.close()

    def test_difference_formula(self, temp_output_path, sample_product_data):
        """差異の数式が正しいかテスト"""
        config = TemplateConfig(num_blocks=1, default_output_path=str(temp_output_path))
        creator = HaibunTemplateCreator(config=config)
        creator.create_template(products=[sample_product_data])

        wb = openpyxl.load_workbook(temp_output_path)
        ws = wb.active

        # AU10（detail_row）に差異の数式があるか
        assert ws['AU10'].value == '=AU9-AT9', "差異の数式が正しくありません"

        wb.close()


if __name__ == '__main__':
    # テストを直接実行する場合
    pytest.main([__file__, '-v'])
